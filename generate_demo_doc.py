import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, KeepTogether
from reportlab.lib import colors

BASE_DIR = r"C:\Users\Nupur\ApniLeap"
DOCX_PATH = os.path.join(BASE_DIR, "Demo_KPIs_and_Challenges.docx")
PDF_PATH = os.path.join(BASE_DIR, "Demo_KPIs_and_Challenges.pdf")

# Data definitions
PROJECTS_DATA = [
    {
        "code": "AL-KLE-026",
        "jira": "ALKLE026",
        "title": "Smart Campus Lab & Resource Optimizer",
        "theme": "Autonomous Robotics & Smart Campus",
        "kpis": [
            {
                "name": "Model Inference & System Latency",
                "target": "50",
                "unit": "ms",
                "sample_measure": "42.8",
                "evidence": "Triton Dynamic Batching benchmark logs on NVIDIA RTX 4090 cluster."
            },
            {
                "name": "Lab Seat & GPU Utilization Efficiency",
                "target": "85",
                "unit": "%",
                "sample_measure": "88.4",
                "evidence": "Campus telemetry MQTT broker aggregated report across 6 computing labs."
            },
            {
                "name": "Pipeline Allocation Throughput",
                "target": "120",
                "unit": "req/sec",
                "sample_measure": "135",
                "evidence": "Locust stress test results showing zero dropped requests under peak load."
            }
        ],
        "challenges": [
            {
                "title": "High GPU memory contention during concurrent model inferencing",
                "root_cause": "Multiple client microservices allocating unshared VRAM buffers without request throttling.",
                "impact": "Inference latency spikes up to 340ms, causing timeouts in campus desktop app.",
                "support_required": "Server access to configure Triton Dynamic Batching with priority queues.",
                "action": "Deploy Triton dynamic queue manager and set maximum batch size to 8.",
                "owner": "KLE Student Lead",
                "expected_jira": "Issue automatically created in Jira board ALKLE026 with Bug/Task type."
            },
            {
                "title": "MQTT sensor heartbeat packets dropped under campus firewall rules",
                "root_cause": "Outbound port 1883 blocked on primary academic VLAN subnet.",
                "impact": "Real-time occupancy status displays 0 seats available in Lab 204.",
                "support_required": "Campus Network Admin to whitelist port 1883 or enable WSS bridge.",
                "action": "Switch sensor telemetry client to Secure WebSockets (port 443/WSS).",
                "owner": "Vihaan Hegde",
                "expected_jira": "Issue automatically synced to Jira board ALKLE026 with student-blocker label."
            }
        ]
    },
    {
        "code": "AL-KLE-023",
        "jira": "ALKLE023",
        "title": "Adaptive Video Streaming & Buffer Management Engine",
        "theme": "Healthcare & Multimedia Systems",
        "kpis": [
            {
                "name": "End-to-End Streaming Latency",
                "target": "50",
                "unit": "ms",
                "sample_measure": "38.5",
                "evidence": "WebRTC peer-to-peer data channel telemetry timestamp difference."
            },
            {
                "name": "Buffer Stall / Rebuffering Ratio",
                "target": "0.5",
                "unit": "%",
                "sample_measure": "0.32",
                "evidence": "ExoPlayer analytics telemetry under synthetic 3G/4G bandwidth throttling."
            },
            {
                "name": "Frame Rendering Throughput",
                "target": "60",
                "unit": "fps",
                "sample_measure": "59.4",
                "evidence": "Hardware-accelerated OpenGL surface view frame counter log."
            }
        ],
        "challenges": [
            {
                "title": "Dataset preprocessing RAM overflow during video chunk indexing",
                "root_cause": "Loading uncompressed 4K raw video chunks directly into worker RAM.",
                "impact": "Training job crashes with Exit Code 137 (OOM Killer) at epoch 4.",
                "support_required": "Allocation of compute VM with PyTorch memory mapping support.",
                "action": "Implement streaming PyTorch chunk generators with disk memory-mapping (mmap).",
                "owner": "Aarav Kulkarni",
                "expected_jira": "Issue automatically created in Jira board ALKLE023."
            },
            {
                "title": "Adaptive bitrate (ABR) oscillations during rapid network handoff",
                "root_cause": "Throughput estimation heuristic reacting aggressively to packet burst drops.",
                "impact": "Frequent quality downgrades causing visual buffering artifacts for viewers.",
                "support_required": "Faculty Mentor guidance on BOLA (Buffer Occupancy based Lyapunov) tuning.",
                "action": "Implement exponential moving average filter on throughput estimation window.",
                "owner": "Rohan Patil",
                "expected_jira": "Issue created and linked with Corrective Action in Jira board ALKLE023."
            }
        ]
    }
]

def generate_pdf():
    doc = SimpleDocTemplate(
        PDF_PATH,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1a3c6e')
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10.5,
        leading=15,
        textColor=colors.HexColor('#555555')
    )
    
    h2_style = ParagraphStyle(
        'H2Style',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=17,
        textColor=colors.HexColor('#1a3c6e'),
        spaceBefore=12,
        spaceAfter=6
    )
    
    h3_style = ParagraphStyle(
        'H3Style',
        parent=styles['Heading3'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#216e4e'),
        spaceBefore=8,
        spaceAfter=4
    )
    
    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#101214')
    )
    
    code_style = ParagraphStyle(
        'CodeStyle',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#0052cc')
    )
    
    badge_style = ParagraphStyle(
        'BadgeStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white
    )
    
    story = []
    
    # Header Title
    story.append(Paragraph("ApniLeap — Demo Data Cheat Sheet", title_style))
    story.append(Paragraph("Pre-Validated Sample <b>KPIs</b> and <b>Challenges</b> for Live Portal & Jira Demonstration", subtitle_style))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#1a3c6e'), spaceAfter=10))
    
    # Instructions Box
    tip_text = (
        "<b>Demo Quick Tip:</b> Copy and paste the values below directly into the ApniLeap web interface. "
        "Adding a <b>Challenge</b> automatically creates a corresponding issue on Jira Cloud Kanban board. "
        "Adding a <b>KPI</b> and recording measurements updates the live completion metrics and project tracking."
    )
    tip_table = Table([[Paragraph(tip_text, body_style)]], colWidths=[540])
    tip_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#eef4fb')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#b3d4fc')),
        ('TOPPADDING', (0,0), (-1,-1), 8),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ('LEFTPADDING', (0,0), (-1,-1), 10),
        ('RIGHTPADDING', (0,0), (-1,-1), 10),
    ]))
    story.append(tip_table)
    story.append(Spacer(1, 10))
    
    for proj in PROJECTS_DATA:
        proj_header = f"Project: {proj['code']} — {proj['title']}  [Jira Board: {proj['jira']}]"
        story.append(Paragraph(proj_header, h2_style))
        
        # --- KPIs Table ---
        story.append(Paragraph(f"1. Ready-to-Input KPIs (Project Tracking &rarr; KPIs Tab &rarr; + Add KPI)", h3_style))
        
        kpi_rows = [
            [
                Paragraph("<b>KPI Name</b>", body_style),
                Paragraph("<b>Target</b>", body_style),
                Paragraph("<b>Unit</b>", body_style),
                Paragraph("<b>Sample Measurement</b>", body_style),
                Paragraph("<b>Evidence Note</b>", body_style)
            ]
        ]
        
        for k in proj['kpis']:
            kpi_rows.append([
                Paragraph(f"<b>{k['name']}</b>", body_style),
                Paragraph(k['target'], code_style),
                Paragraph(k['unit'], code_style),
                Paragraph(f"<b>{k['sample_measure']} {k['unit']}</b>", body_style),
                Paragraph(k['evidence'], body_style)
            ])
            
        kpi_table = Table(kpi_rows, colWidths=[130, 45, 45, 80, 240])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#e6f4ea')),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#d9dee5')),
            ('TOPPADDING', (0,0), (-1,-1), 5),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
            ('LEFTPADDING', (0,0), (-1,-1), 6),
            ('RIGHTPADDING', (0,0), (-1,-1), 6),
        ]))
        story.append(kpi_table)
        story.append(Spacer(1, 10))
        
        # --- Challenges Table ---
        story.append(Paragraph(f"2. Ready-to-Input Challenges (Project Tracking &rarr; Challenges Tab &rarr; + Add Challenge)", h3_style))
        
        ch_rows = [
            [
                Paragraph("<b>Field</b>", body_style),
                Paragraph("<b>Challenge #1 (Demo Copy-Paste)</b>", body_style),
                Paragraph("<b>Challenge #2 (Demo Copy-Paste)</b>", body_style)
            ]
        ]
        
        ch1 = proj['challenges'][0]
        ch2 = proj['challenges'][1]
        
        fields = [
            ("Title", ch1['title'], ch2['title']),
            ("Root Cause", ch1['root_cause'], ch2['root_cause']),
            ("Impact", ch1['impact'], ch2['impact']),
            ("Support Required", ch1['support_required'], ch2['support_required']),
            ("Action Description", ch1['action'], ch2['action']),
            ("Action Owner", ch1['owner'], ch2['owner']),
            ("Jira Sync Effect", ch1['expected_jira'], ch2['expected_jira']),
        ]
        
        for fname, val1, val2 in fields:
            ch_rows.append([
                Paragraph(f"<b>{fname}</b>", body_style),
                Paragraph(val1, body_style),
                Paragraph(val2, body_style),
            ])
            
        ch_table = Table(ch_rows, colWidths=[100, 220, 220])
        ch_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#fff6df')),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#d9dee5')),
            ('TOPPADDING', (0,0), (-1,-1), 5),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
            ('LEFTPADDING', (0,0), (-1,-1), 6),
            ('RIGHTPADDING', (0,0), (-1,-1), 6),
        ]))
        story.append(ch_table)
        story.append(Spacer(1, 14))
        
    doc.build(story)
    print("PDF generated successfully:", PDF_PATH)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def generate_docx():
    doc = Document()
    
    # Page setup - 0.5 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        
    # Title
    title = doc.add_paragraph()
    r = title.add_run("ApniLeap — Demo Data Cheat Sheet")
    r.font.name = "Arial"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = RGBColor(26, 60, 110) # #1a3c6e
    title.paragraph_format.space_after = Pt(2)
    
    sub = doc.add_paragraph()
    r_sub = sub.add_run("Pre-Validated Sample KPIs and Challenges for Live Demonstration & Jira Sync")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(85, 85, 85)
    sub.paragraph_format.space_after = Pt(10)
    
    # Tip Box
    box_p = doc.add_paragraph()
    box_p.paragraph_format.left_indent = Inches(0.1)
    r_tip = box_p.add_run("Quick Demo Note: ")
    r_tip.bold = True
    r_tip.font.color.rgb = RGBColor(10, 77, 176)
    r_text = box_p.add_run(
        "Copy and paste these exact inputs during your presentation. When you add a Challenge in ApniLeap, "
        "it immediately triggers a REST API call to Jira Cloud creating the corresponding issue in the Kanban board. "
        "When you add a KPI, you can record sample measurements to showcase student tracking."
    )
    r_text.font.size = Pt(9.5)
    box_p.paragraph_format.space_after = Pt(14)
    
    for proj in PROJECTS_DATA:
        # Project Heading
        h2 = doc.add_paragraph()
        h2.paragraph_format.space_before = Pt(12)
        h2.paragraph_format.space_after = Pt(4)
        r_h2 = h2.add_run(f"Project: {proj['code']} — {proj['title']}  (Jira: {proj['jira']})")
        r_h2.font.name = "Arial"
        r_h2.font.size = Pt(13)
        r_h2.bold = True
        r_h2.font.color.rgb = RGBColor(26, 60, 110)
        
        # 1. KPIs
        h3 = doc.add_paragraph()
        h3.paragraph_format.space_before = Pt(6)
        h3.paragraph_format.space_after = Pt(4)
        r_h3 = h3.add_run("1. Demo KPIs (Project Dashboard -> KPIs Tab -> + Add KPI / Add Measurement)")
        r_h3.font.name = "Arial"
        r_h3.font.size = Pt(11)
        r_h3.bold = True
        r_h3.font.color.rgb = RGBColor(30, 126, 52)
        
        # KPI Table
        t_kpi = doc.add_table(rows=1, cols=5)
        t_kpi.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = t_kpi.rows[0].cells
        hdr_titles = ["KPI Name", "Target", "Unit", "Sample Measure", "Evidence Note"]
        for idx, htitle in enumerate(hdr_titles):
            hdr_cells[idx].text = htitle
            hdr_cells[idx].paragraphs[0].runs[0].font.bold = True
            hdr_cells[idx].paragraphs[0].runs[0].font.size = Pt(9)
            set_cell_background(hdr_cells[idx], "E6F4EA")
            
        for k in proj['kpis']:
            row_cells = t_kpi.add_row().cells
            row_cells[0].text = k['name']
            row_cells[1].text = k['target']
            row_cells[2].text = k['unit']
            row_cells[3].text = f"{k['sample_measure']} {k['unit']}"
            row_cells[4].text = k['evidence']
            for c in row_cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8.5)
                        
        doc.add_paragraph().paragraph_format.space_after = Pt(8)
        
        # 2. Challenges
        h3_c = doc.add_paragraph()
        h3_c.paragraph_format.space_before = Pt(6)
        h3_c.paragraph_format.space_after = Pt(4)
        r_h3c = h3_c.add_run("2. Demo Challenges (Project Dashboard -> Challenges Tab -> + Add Challenge)")
        r_h3c.font.name = "Arial"
        r_h3c.font.size = Pt(11)
        r_h3c.bold = True
        r_h3c.font.color.rgb = RGBColor(122, 86, 0)
        
        t_ch = doc.add_table(rows=1, cols=3)
        t_ch.alignment = WD_TABLE_ALIGNMENT.CENTER
        ch_hdr = t_ch.rows[0].cells
        ch_hdr[0].text = "Modal Field"
        ch_hdr[1].text = "Challenge #1 (Copy & Paste)"
        ch_hdr[2].text = "Challenge #2 (Copy & Paste)"
        for c in ch_hdr:
            c.paragraphs[0].runs[0].font.bold = True
            c.paragraphs[0].runs[0].font.size = Pt(9)
            set_cell_background(c, "FFF6DF")
            
        ch1 = proj['challenges'][0]
        ch2 = proj['challenges'][1]
        fields = [
            ("Title", ch1['title'], ch2['title']),
            ("Root Cause", ch1['root_cause'], ch2['root_cause']),
            ("Impact", ch1['impact'], ch2['impact']),
            ("Support Required", ch1['support_required'], ch2['support_required']),
            ("Corrective Action", ch1['action'], ch2['action']),
            ("Action Owner", ch1['owner'], ch2['owner']),
            ("Jira Cloud Sync", ch1['expected_jira'], ch2['expected_jira']),
        ]
        
        for fname, val1, val2 in fields:
            row_cells = t_ch.add_row().cells
            row_cells[0].text = fname
            row_cells[0].paragraphs[0].runs[0].font.bold = True
            row_cells[1].text = val1
            row_cells[2].text = val2
            for c in row_cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8.5)
                        
        doc.add_paragraph().paragraph_format.space_after = Pt(14)
        
    doc.save(DOCX_PATH)
    print("DOCX generated successfully:", DOCX_PATH)

if __name__ == "__main__":
    generate_pdf()
    generate_docx()
